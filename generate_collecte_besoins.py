"""Generate the 'Collecte des besoins en formations P4' Excel workbook.

Sheets: INFO, ASI, BD, GL, IL, Telecom, Tronc Commun.
Structure:
    Row 1: Title (merged across all 11 columns)
    Row 2: blank
    Row 3: Column headers (salmon background)
    Row 4+: Data rows (alternating colors, vertical UP merges,
            dropdown validation on column H)
    Empty rows left blank below the sample data for new entries.
"""

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.page import PageMargins
from openpyxl.worksheet.views import Pane


# --- Constants ----------------------------------------------------------------
OUTPUT_FILE = "collecte_besoins_formations_P4.xlsx"

UP_SHEETS = ["ASI", "BD", "GL", "IL", "Telecom", "Tronc Commun"]
INFO_SHEET = "INFO"

TITLE_TEXT = "Collecte des besoins en formations P4"

HEADERS = [
    "Département UP",
    "",
    "Besoin en formation",
    "Objectifs de la formation et prérequis",
    "Proposition d'un formateur (si c'est possible)",
    "Volume horaire souhaité",
    "Dates et créneaux souhaités",
    "Type: ouverte/fermée",
    (
        "Nombre de participants maximal (dans le cas ou elle sera ouverte "
        "à d'autres UPs)"
    ),
    "Liste des participants de l'UP",
    "Autres informations utiles à ajouter",
]

# Sample data per UP sheet (rows). Empty lists => only the title/header is
# written, the user can fill in fresh rows below.
SAMPLE_DATA: dict[str, list[list[str]]] = {
    "ASI": [
        [
            "ASI",
            "",
            "MLOps",
            "Préparer les enseignants à enseigner le module",
            "Abderrahmen BEN AROUS + Ghassen HAMMOUDA",
            "18h",
            "Les mercredis apm du mai",
            "Fermée",
            "",
            (
                "Nadine MILI, Ons BEN SALAH, Leila BEN DHIEF, Alaa RAMI, "
                "Ichraf AYARI, Erij KLAI, Lassaad SAIDAINI, Mohamed RJAB, "
                "Nawres MRABET"
            ),
            "",
        ],
    ],
    "BD": [
        [
            "BD",
            "",
            "NoSQL",
            "Préparer les enseignants à enseigner le module",
            "Lilia Gossa",
            "21h",
            "Les mercredis apm du avril et mai",
            "Ouverte",
            "",
            "",
            "",
        ],
        [
            "",
            "",
            "Machine Learning",
            "Préparer les enseignants à enseigner le module",
            "Oumayma Guessmi, Dorsaf Hrizi (6h)",
            "21h",
            "Les mercredis apm du avril et mai",
            "Fermée",
            "",
            "UP-BD + UP-IL",
            "",
        ],
        [
            "",
            "",
            "Machine Learning avancé",
            "",
            (
                "Jihene Hlel (6h), Jihene Hlel / Ines Slimen / Sonia Mesbah "
                "(selon disponibilité)"
            ),
            "12h",
            "Les mercredis apm du avril",
            "Fermée",
            "",
            "UP-BD",
            "",
        ],
    ],
    "GL": [],
    "IL": [
        [
            "IL",
            "",
            "Deep Learning",
            "",
            "",
            "30h",
            "Les mercredis apm du Mai + Juin",
            "Fermée",
            "",
            "UP IL",
            "",
        ],
    ],
    "Telecom": [],
    "Tronc Commun": [],
}

# --- Style helpers ------------------------------------------------------------
SALMON_HEX = "F4CCCC"
LIGHT_GRAY_HEX = "F2F2F2"
INFO_TITLE_HEX = "2E5984"
INFO_BODY_HEX = "DEEBF7"

THIN = Side(style="thin", color="808080")
THIN_BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

HEADER_FILL = PatternFill("solid", fgColor=SALMON_HEX)
ALT_FILL = PatternFill("solid", fgColor=LIGHT_GRAY_HEX)
INFO_TITLE_FILL = PatternFill("solid", fgColor=INFO_TITLE_HEX)
INFO_BODY_FILL = PatternFill("solid", fgColor=INFO_BODY_HEX)

TITLE_FONT = Font(name="Calibri", size=18, bold=True, color="FFFFFF")
HEADER_FONT = Font(name="Calibri", size=11, bold=True, color="000000")
DATA_FONT = Font(name="Calibri", size=11, color="000000")
INFO_LABEL_FONT = Font(name="Calibri", size=11, bold=True, color="FFFFFF")

CENTER_WRAP = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT_WRAP = Alignment(horizontal="left", vertical="center", wrap_text=True)
RIGHT_WRAP = Alignment(horizontal="right", vertical="center", wrap_text=True)


# --- Sheet builders -----------------------------------------------------------
def style_header_row(ws, row: int) -> None:
    """Apply the salmon header style to the merged A3:B3 + C3..K3 row."""
    for col in range(1, 12):
        cell = ws.cell(row=row, column=col)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = CENTER_WRAP
        cell.border = THIN_BORDER


def write_title(ws) -> None:
    """Row 1: merged title across A1:K1."""
    ws.merge_cells("A1:K1")
    title = ws["A1"]
    title.value = TITLE_TEXT
    title.font = TITLE_FONT
    title.alignment = CENTER_WRAP
    title.fill = INFO_TITLE_FILL
    ws.row_dimensions[1].height = 34
    ws.row_dimensions[2].height = 6  # spacer


def write_headers(ws) -> None:
    """Row 3: merged UP column + individual column headers."""
    # Merge the UP label across A:B
    ws.merge_cells(start_row=3, start_column=1, end_row=3, end_column=2)
    ws.cell(row=3, column=1).value = "Département UP"

    for idx, header in enumerate(HEADERS[2:], start=3):  # skip the empty placeholder
        ws.cell(row=3, column=idx).value = header

    style_header_row(ws, 3)
    ws.row_dimensions[3].height = 56


def write_data_rows(ws, rows: list[list[str]]) -> int:
    """Write the sample rows starting at row 4. Returns last used row."""
    start = 4
    for r_idx, row in enumerate(rows):
        excel_row = start + r_idx
        for c_idx, value in enumerate(row, start=1):
            cell = ws.cell(row=excel_row, column=c_idx)
            cell.value = value
            cell.font = DATA_FONT
            cell.border = THIN_BORDER
            if r_idx % 2 == 1:
                cell.fill = ALT_FILL
            # Default alignment: long text wrapped left-aligned
            if c_idx == 6:  # Volume horaire
                cell.alignment = RIGHT_WRAP
            elif c_idx == 8:  # Type
                cell.alignment = CENTER_WRAP
            else:
                cell.alignment = LEFT_WRAP
        # Tall row for wrapped content
        ws.row_dimensions[excel_row].height = 48
    return start + len(rows) - 1


def merge_up_column(ws, rows: list[list[str]], start: int) -> None:
    """Vertically merge cells in column A for consecutive identical UPs."""
    if not rows:
        return
    # Build (excel_row, up_label) pairs (only the first row of a run carries
    # the label; subsequent rows have an empty string in column A).
    first_label_row = None
    current_label = None
    for offset, row in enumerate(rows):
        excel_row = start + offset
        label = row[0]
        if label:
            if current_label is not None and first_label_row is not None and excel_row - 1 > first_label_row:
                ws.merge_cells(
                    start_row=first_label_row, start_column=1,
                    end_row=excel_row - 1, end_column=2,
                )
            current_label = label
            first_label_row = excel_row
        # else: continuation of the same run, nothing to do
    if current_label and first_label_row and start + len(rows) - 1 > first_label_row:
        ws.merge_cells(
            start_row=first_label_row, start_column=1,
            end_row=start + len(rows) - 1, end_column=2,
        )


def add_type_validation(ws, first_row: int, last_row: int) -> None:
    """Add a data-validation dropdown on column H (Type) for all data rows."""
    dv = DataValidation(
        type="list",
        formula1='"Ouverte,Fermée"',
        allow_blank=True,
        showDropDown=False,  # show the dropdown arrow
    )
    dv.error = "Veuillez choisir « Ouverte » ou « Fermée »."
    dv.errorTitle = "Valeur invalide"
    dv.prompt = "Choisir Ouverte ou Fermée"
    dv.promptTitle = "Type de formation"
    dv.add(f"H{first_row}:H{last_row}")
    ws.add_data_validation(dv)


def configure_view_and_print(ws, last_data_row: int) -> None:
    """Set column widths, freeze pane, auto-filter, print setup."""
    widths = {
        "A": 14, "B": 4,
        "C": 20, "D": 35, "E": 30, "F": 12, "G": 25,
        "H": 12, "I": 20, "J": 30, "K": 25,
    }
    for col, width in widths.items():
        ws.column_dimensions[col].width = width

    # Freeze top rows so the header (row 3) stays visible
    ws.freeze_panes = "A4"

    # Auto-filter on the header row
    ws.auto_filter.ref = f"A3:K{last_data_row}"

    # Print setup
    ws.page_setup.orientation = ws.ORIENTATION_LANDSCAPE
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_margins = PageMargins(left=0.5, right=0.5, top=0.7, bottom=0.7)
    ws.print_options.horizontalCentered = True
    ws.print_title_rows = "1:3"


# --- INFO sheet ---------------------------------------------------------------
INFO_SECTIONS = [
    ("Contexte", [
        "Cette fiche recense les besoins en formation des différentes Unités Pédagogiques (UP) "
        "pour le semestre P4.",
        "Merci de compléter une ligne par besoin identifié.",
    ]),
    ("Comment remplir la fiche ?", [
        "1. Choisir l'UP concernée (colonne A).",
        "2. Renseigner l'intitulé exact du besoin (colonne C).",
        "3. Décrire les objectifs et prérequis (colonne D).",
        "4. Proposer un formateur interne ou externe si possible (colonne E).",
        "5. Indiquer le volume horaire, les dates et créneaux souhaités.",
        "6. Préciser si la formation est ouverte à d'autres UPs (colonne H).",
        "7. Lister les participants pressentis (colonne J).",
    ]),
    ("Légende — Type de formation", [
        "Ouverte : accessible à d'autres UPs (renseigner la colonne I).",
        "Fermée : réservée aux enseignants de l'UP (renseigner la colonne J).",
    ]),
    ("Contact", [
        "Pour toute question, contacter le responsable formation de votre UP.",
    ]),
]


def build_info_sheet(ws) -> None:
    """Build the INFO sheet with usage instructions."""
    ws.merge_cells("A1:K1")
    title = ws["A1"]
    title.value = TITLE_TEXT
    title.font = Font(name="Calibri", size=20, bold=True, color="FFFFFF")
    title.alignment = CENTER_WRAP
    title.fill = INFO_TITLE_FILL
    ws.row_dimensions[1].height = 38

    row = 3
    for section_title, lines in INFO_SECTIONS:
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=11)
        cell = ws.cell(row=row, column=1)
        cell.value = section_title
        cell.font = INFO_LABEL_FONT
        cell.alignment = Alignment(horizontal="left", vertical="center", indent=1)
        cell.fill = PatternFill("solid", fgColor=INFO_TITLE_HEX)
        ws.row_dimensions[row].height = 24
        row += 1
        for line in lines:
            ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=11)
            cell = ws.cell(row=row, column=1)
            cell.value = line
            cell.font = DATA_FONT
            cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True, indent=1)
            cell.fill = INFO_BODY_FILL
            ws.row_dimensions[row].height = 22
            row += 1
        row += 1  # spacer

    # Column widths for the info sheet
    for col, width in {"A": 18, "B": 18, "C": 22, "D": 35, "E": 30, "F": 12, "G": 25, "H": 12, "I": 20, "J": 30, "K": 25}.items():
        ws.column_dimensions[col].width = width


# --- Main ---------------------------------------------------------------------
def main() -> None:
    wb = Workbook()

    # The default sheet becomes INFO.
    info_ws = wb.active
    info_ws.title = INFO_SHEET
    build_info_sheet(info_ws)

    for up in UP_SHEETS:
        ws = wb.create_sheet(title=up)
        write_title(ws)
        write_headers(ws)
        rows = SAMPLE_DATA.get(up, [])
        last_row = write_data_rows(ws, rows)
        merge_up_column(ws, rows, start=4)
        # Validation + filter + freeze + print apply to the whole table area
        # (header at row 3, sample data ends at last_row). We extend to row 60
        # so future entries are also covered.
        last_relevant_row = max(last_row, 60)
        add_type_validation(ws, first_row=4, last_row=last_relevant_row)
        configure_view_and_print(ws, last_data_row=last_relevant_row)

    wb.save(OUTPUT_FILE)
    print(f"Fichier généré : {OUTPUT_FILE}")


if __name__ == "__main__":
    main()
