"""Génération de rapports analytiques exportables (Excel / PDF).

Équivalents Python de POI/iText (interdits ici : ce service est en FastAPI) :
- Excel via openpyxl
- PDF via reportlab
Les bibliothèques sont importées paresseusement : un export reste indisponible
(501) plutôt que de faire échouer le démarrage si la dépendance manque.
"""

import io
import logging
from datetime import date
from typing import Any

from fastapi import HTTPException
from sqlalchemy.orm import Session

from app.config import settings
from app.engines.reporting_engine import ReportingEngine

logger = logging.getLogger(__name__)

_COL_DEPARTEMENT = "Département"

# (type d'export) → (titre feuille, en-têtes colonnes, extracteur de lignes)
_EXCEL_COLUMNS = {
    "INACTIFS": (
        "Enseignants inactifs",
        ["Nom", "Prénom", "Email", _COL_DEPARTEMENT, "UP",
         "Dernière formation", "Mois sans formation", "Score risque", "Niveau"],
        lambda r: [r["nom"], r["prenom"], r["email"], r.get("departement"), r.get("up"),
                   r.get("derniereFormationDate"), r["nombreMoisDepuisDerniereFormation"],
                   r["scoreRisqueDecrochage"], r["niveauRisque"]],
    ),
    "PAR_UP": (
        "Analyse par UP",
        ["UP", _COL_DEPARTEMENT, "Enseignants", "Formations", "Participations",
         "Taux participation %", "Inactifs", "Score engagement"],
        lambda r: [r["upNom"], r.get("departementNom"), r["nombreEnseignants"],
                   r["nombreFormationsOrganisees"], r["nombreParticipations"],
                   r["tauxParticipation"], r["enseignantsSansFormation"], r["scoreEngagement"]],
    ),
    "PAR_DEPT": (
        "Analyse par département",
        [_COL_DEPARTEMENT, "Enseignants", "Formations", "Participations",
         "Taux participation %", "% à risque", "Niveau moyen", "Score engagement"],
        lambda r: [r["departementNom"], r["nombreEnseignants"], r["nombreFormationsOrganisees"],
                   r["nombreParticipations"], r["tauxParticipation"], r["pourcentageARisque"],
                   r["niveauCompetenceMoyen"], r["scoreEngagement"]],
    ),
}


def _rows_for(engine: ReportingEngine, type_export: str, *,
              mois: int, annee: int | None, departement: str | None, up: str | None) -> list[dict[str, Any]]:
    if type_export == "INACTIFS":
        return engine.enseignants_sans_formation(
            mois, departement, up, page=0, size=settings.export_max_rows,
        )["items"]
    if type_export == "PAR_UP":
        return engine.formations_par_up(annee, departement)
    if type_export == "PAR_DEPT":
        return engine.formations_par_departement(annee)["departements"]
    raise HTTPException(status_code=400, detail=f"Type d'export inconnu: {type_export}")


def build_excel(db: Session, type_export: str, *, mois: int, annee: int | None,
                departement: str | None, up: str | None) -> tuple[bytes, str]:
    if type_export not in _EXCEL_COLUMNS:
        raise HTTPException(status_code=400, detail=f"Type d'export inconnu: {type_export}")
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
    except ImportError as exc:  # pragma: no cover
        raise HTTPException(status_code=501, detail="openpyxl non installé sur le service.") from exc

    title, headers, extract = _EXCEL_COLUMNS[type_export]
    rows = _rows_for(ReportingEngine(db), type_export, mois=mois, annee=annee,
                     departement=departement, up=up)

    wb = Workbook()
    ws = wb.active
    ws.title = title[:31]
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="1B5E20")
    for col, h in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.font = header_font
        cell.fill = header_fill
    for ri, r in enumerate(rows, start=2):
        for ci, val in enumerate(extract(r), start=1):
            ws.cell(row=ri, column=ci, value=val)
    for col in range(1, len(headers) + 1):
        ws.column_dimensions[chr(64 + col)].width = 20

    buf = io.BytesIO()
    wb.save(buf)
    filename = f"rapport_{type_export.lower()}_{date.today().isoformat()}.xlsx"
    return buf.getvalue(), filename


def build_pdf(db: Session, type_rapport: str, *, annee: int | None) -> tuple[bytes, str]:
    try:
        from reportlab.lib import colors
        from reportlab.lib.pagesizes import A4
        from reportlab.lib.units import cm
        from reportlab.platypus import (SimpleDocTemplate, Table, TableStyle,
                                        Paragraph, Spacer)
        from reportlab.lib.styles import getSampleStyleSheet
    except ImportError as exc:  # pragma: no cover
        raise HTTPException(status_code=501, detail="reportlab non installé sur le service.") from exc

    engine = ReportingEngine(db)
    annee_eff = annee or date.today().year
    dept_data = engine.formations_par_departement(annee_eff)["departements"]

    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, topMargin=2 * cm, bottomMargin=2 * cm)
    styles = getSampleStyleSheet()
    titre = "Rapport mensuel" if type_rapport == "RAPPORT_MENSUEL" else "Rapport annuel"
    story = [
        Paragraph(f"D2F — Analyse Prédictive — {titre} {annee_eff}", styles["Title"]),
        Spacer(1, 0.5 * cm),
        Paragraph("Synthèse par département", styles["Heading2"]),
        Spacer(1, 0.3 * cm),
    ]
    table_data = [["Département", "Enseignants", "Formations", "Taux part. %", "% à risque", "Engagement"]]
    for d in dept_data:
        table_data.append([
            d["departementNom"], d["nombreEnseignants"], d["nombreFormationsOrganisees"],
            d["tauxParticipation"], d["pourcentageARisque"], d["scoreEngagement"],
        ])
    table = Table(table_data, repeatRows=1)
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1B5E20")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F1F8E9")]),
        ("FONTSIZE", (0, 0), (-1, -1), 9),
    ]))
    story.append(table)
    doc.build(story)
    filename = f"rapport_{type_rapport.lower()}_{annee_eff}.pdf"
    return buf.getvalue(), filename
